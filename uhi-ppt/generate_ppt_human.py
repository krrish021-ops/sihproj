"""HeatSight UHI hackathon PPT — HUMAN edition.
Plain white slides, Calibri, simple bullets/tables, honest 'proposed' framing,
source footnotes. No dashboard mockup, no glow art, no ghost numerals.
Output: uhi-ppt/HeatSight_UHI_PPT_Human.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

INK    = RGBColor(0x21, 0x21, 0x21)
GRAY   = RGBColor(0x61, 0x61, 0x61)
LGRAY  = RGBColor(0x9E, 0x9E, 0x9E)
LINE   = RGBColor(0xE0, 0xE0, 0xE0)
BOX    = RGBColor(0xF5, 0xF5, 0xF5)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
RED    = RGBColor(0xA9, 0x32, 0x26)   # single brick-red accent, used sparingly
BLUE   = RGBColor(0x2E, 0x5E, 0x8A)   # chart second series
FONT   = "Calibri"
MONO   = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.33

def shape_rect(s, l, t, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    return sh

def txt(s, l, t, w, h, blocks, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """blocks: list of paragraphs; paragraph = list of (text, size, color, bold, font)."""
    tx = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tx.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.02); tf.margin_right = Inches(0.02)
    first = True
    for para in blocks:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(4)
        for piece in para:
            text, size, color, bold, font = (list(piece) + [FONT])[:5]
            r = p.add_run(); r.text = text
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.bold = bool(bold); r.font.name = font or FONT
    return tx

def footer(s, num):
    shape_rect(s, 0.8, 7.0, W - 1.6, 0.015, fill=LINE)
    txt(s, 0.8, 7.06, 8.0, 0.3, [[("HeatSight — Urban Heat Island Forecasting  |  SIH 2026", 9, LGRAY, False, None)]])
    txt(s, 11.9, 7.06, 0.63, 0.3, [[(str(num), 9, LGRAY, False, None)]], align=PP_ALIGN.RIGHT)

def header(s, num, title, subtitle=None):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    txt(s, 0.8, 0.45, 11.7, 0.6, [[(title, 28, INK, True, None)]])
    if subtitle:
        txt(s, 0.8, 1.05, 11.7, 0.4, [[(subtitle, 13, GRAY, False, None)]])
    y = 1.55 if subtitle else 1.25
    shape_rect(s, 0.8, y, W - 1.6, 0.018, fill=LINE)
    shape_rect(s, 0.8, y, 1.3, 0.035, fill=RED)
    footer(s, num)
    return y + 0.3  # content top

def bullets(s, l, t, w, items, size=14, gap=10):
    """items: (bold_lead, rest) or plain str."""
    tx = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(7.0 - t))
    tf = tx.text_frame; tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap); p.space_before = Pt(2)
        r0 = p.add_run(); r0.text = "•  "
        r0.font.size = Pt(size); r0.font.color.rgb = RED; r0.font.name = FONT
        if isinstance(item, tuple):
            lead, rest = item
            r = p.add_run(); r.text = lead
            r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = INK; r.font.name = FONT
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(size); r2.font.color.rgb = INK; r2.font.name = FONT
        else:
            r = p.add_run(); r.text = item
            r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = FONT

def simple_table(s, l, t, w, data, widths, size=11.5, row_h=0.55, first_col_bold=True):
    """Plain PowerPoint-style table: gray header, thin dividers."""
    gt = s.shapes.add_table(len(data), len(data[0]), Inches(l), Inches(t), Inches(w), Inches(row_h * len(data))).table
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            c = gt.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xE8, 0xE8, 0xE8) if i == 0 else WHITE
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = Inches(0.1); c.margin_right = Inches(0.08)
            c.margin_top = Inches(0.04); c.margin_bottom = Inches(0.04)
            p = c.text_frame.paragraphs[0]
            p.space_after = Pt(1)
            r = p.add_run(); r.text = val
            r.font.name = FONT
            r.font.size = Pt(size if i > 0 else size + 0.5)
            r.font.bold = (i == 0) or (j == 0 and first_col_bold)
            r.font.color.rgb = INK
    for j, wd in enumerate(widths):
        gt.columns[j].width = Inches(wd)
    # thin dividers under each row (looks like a formatted PPT table)
    for i in range(len(data)):
        yy = t + row_h * (i + 1)
        shape_rect(s, l, yy - 0.015, w, 0.015, fill=LINE)
    shape_rect(s, l, t - 0.01, w, 0.03, fill=RED)  # small red top edge
    return t + row_h * len(data)

def note_box(s, l, t, w, h, text, size=11.5):
    shape_rect(s, l, t, w, h, fill=BOX, line=LINE, lw=0.75)
    shape_rect(s, l, t, 0.07, h, fill=RED)
    txt(s, l + 0.22, t + 0.08, w - 0.4, h - 0.16, [[(text, size, GRAY, False, None)]], anchor=MSO_ANCHOR.MIDDLE)

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text

# ============ 1 · TITLE ============
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
shape_rect(s, 0, 0, W, 0.12, fill=RED)
txt(s, 0.9, 1.0, 11.5, 0.4, [[("Smart India Hackathon 2026  •  Idea Presentation", 14, GRAY, False, None)]])
txt(s, 0.9, 1.5, 11.5, 1.0, [[("HeatSight: Forecasting Urban Heat Islands", 42, INK, True, None)]])
txt(s, 0.9, 2.5, 11.0, 0.8, [[("Using satellite LST time-series and land-use data to predict heat intensification in growing Indian cities", 17, GRAY, False, None)]])
shape_rect(s, 0.9, 3.5, 1.2, 0.045, fill=RED)
shape_rect(s, 0.9, 3.72, 7.6, 1.55, fill=BOX, line=LINE, lw=0.75)
shape_rect(s, 0.9, 3.72, 0.07, 1.55, fill=RED)
txt(s, 1.25, 3.88, 7.1, 1.3, [
    [("Team: ", 14, INK, True, None), ("[ Your Team Name ]", 14, INK, False, None)],
    [("College: ", 14, INK, True, None), ("[ College Name, City ]", 14, INK, False, None)],
    [("Problem Statement ID: ", 14, INK, True, None), ("[ PS-ID ]", 14, INK, False, None)],
])
txt(s, 0.9, 5.9, 11.5, 0.4, [[("Team members: [ Name 1, Name 2, Name 3, Name 4, Name 5, Name 6 ]", 12, GRAY, False, None)]])
footer(s, 1)
notes(s, "Introduce the team and the one-line idea: predict where Indian cities will overheat next, using free satellite data, so planners can act early.")

# ============ 2 · PROBLEM ============
s = prs.slides.add_slide(BLANK)
top = header(s, 2, "Problem we are solving", "Indian cities are heating up faster than their surroundings")
bullets(s, 0.8, top, 11.7, [
    ("Urban Heat Island effect: ", "built-up city areas stay 3–7 °C hotter than nearby rural areas, even at night."),
    ("Heatwaves are getting worse: ", "Delhi recorded 50 °C+ days; heat-stroke cases rise every summer across North and Central India."),
    ("Why it is happening: ", "concrete and asphalt replacing trees, lakes and open soil — heat gets absorbed in the day and released at night."),
    ("Who suffers: ", "outdoor workers, elderly people, low-income households without cooling, and city power grids during peak demand."),
    ("The planning gap: ", "today no planner can answer — which wards of my city will overheat in the next 5 years?"),
], size=13.5, gap=9)
txt(s, 0.8, 6.35, 11.7, 0.3, [[("Sources: IMD heatwave reports 2023–25; IPCC AR6 (urban warming); news reports on Delhi / Ahmedabad heat action plans.", 9.5, LGRAY, False, None)]])
notes(s, "Explain UHI in one line, then the human cost. End with the planning gap question — that is the problem we solve.")

# ============ 3 · WHY CURRENT SOLUTIONS FALL SHORT ============
s = prs.slides.add_slide(BLANK)
top = header(s, 3, "Why current solutions are not enough", "We studied what exists — each option misses something important")
data = [
    ("What exists today", "Limitation"),
    ("IMD forecasts and weather apps", "Give city-average air temperature. No ward-level detail, no land-surface heat data."),
    ("Research papers with LST maps", "One-time static maps of a city. No future prediction, no usable tool for planners."),
    ("Smart city dashboards", "Track air quality, traffic, etc. Heat-island forecasting is not included anywhere."),
]
simple_table(s, 0.8, top + 0.1, 11.73, data, [3.4, 8.33], size=12.5, row_h=0.72, first_col_bold=True)
note_box(s, 0.8, 5.55, 11.73, 0.95,
         "Our finding: there is no free, open tool that combines satellite heat data + land-use change to forecast future hotspots for planners. That is the gap HeatSight fills.")
notes(s, "Show we did our homework: three things exist, each has a clear limitation. The note box is our conclusion.")

# ============ 4 · OUR PROPOSED SOLUTION ============
s = prs.slides.add_slide(BLANK)
top = header(s, 4, "Our proposed solution — HeatSight", "A forecasting system that tells planners where heat will rise next")
items = [
    ("1. Study the past (2015–2025). ", "Build a 10-year record of land-surface temperature and land-use change for every ward, from free satellite data."),
    ("2. Predict the hotspots (2030). ", "Use machine learning to forecast which wards will warm the fastest as construction and concrete cover grow."),
    ("3. Support planning decisions. ", "Provide a ward heat-risk ranking with practical suggestions — plantations, cool roofs, lake revival, building norms."),
]
bullets_plain = []
tx = s.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.7), Inches(3.6))
tf = tx.text_frame; tf.word_wrap = True
first = True
for lead, rest in items:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(12)
    r = p.add_run(); r.text = lead
    r.font.size = Pt(13.5); r.font.bold = True; r.font.color.rgb = RED; r.font.name = FONT
    r2 = p.add_run(); r2.text = rest
    r2.font.size = Pt(13.5); r2.font.color.rgb = INK; r2.font.name = FONT
note_box(s, 0.8, 5.55, 11.73, 0.95,
         "In one line: past satellite heat data + land-use change → predicted future hotspots → better planning decisions.")
notes(s, "Three steps: past, future, action. Read the one-line summary at the end — that is what judges will remember.")

# ============ 5 · SYSTEM DESIGN ============
s = prs.slides.add_slide(BLANK)
top = header(s, 5, "Proposed system design", "Five simple stages — from satellite data to ward-level forecast")
steps = [
    ("Step 1\nData", "Landsat 8/9\nSentinel-2\nMODIS, IMD"),
    ("Step 2\nClean + Prepare", "Cloud removal\nLST calculation\nCut by ward map"),
    ("Step 3\nFeatures", "LST trend\nNDVI, NDBI\nBuilt-up %"),
    ("Step 4\nML Model", "Random Forest /\nXGBoost\nTrain on 2015–25"),
    ("Step 5\nOutput", "2030 hotspot map\nWard ranking\nAction points"),
]
x = 0.8
bw, bh = 1.95, 1.9
for i, (title, desc) in enumerate(steps):
    shape_rect(s, x, top + 0.15, bw, bh, fill=BOX if i % 2 else WHITE, line=LGRAY, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, x + 0.1, top + 0.28, bw - 0.2, 0.6, [[(title.replace("\n", " — "), 11.5, RED, True, None)]], align=PP_ALIGN.CENTER)
    txt(s, x + 0.1, top + 0.85, bw - 0.2, 1.1, [[(desc, 10.5, GRAY, False, None)]], align=PP_ALIGN.CENTER)
    if i < 4:
        txt(s, x + bw, top + 0.75, 0.5, 0.5, [[("→", 22, LGRAY, True, None)]], align=PP_ALIGN.CENTER)
    x += bw + 0.5
txt(s, 0.8, top + 2.4, 11.73, 0.9, [
    [("Key formulas we will use:  ", 12, INK, True, None)],
    [("NDVI = (NIR − Red)/(NIR + Red)      NDBI = (SWIR − NIR)/(SWIR + NIR)      UHI intensity = LST(urban) − LST(rural)", 11.5, GRAY, False, MONO)],
])
note_box(s, 0.8, 5.6, 11.73, 0.9, "Detailed step-by-step LST calculation is given in Appendix A (last slide).")
notes(s, "Walk through the five boxes left to right. Mention the three formulas briefly and point to the appendix for detail.")

# ============ 6 · DATA AND TOOLS ============
s = prs.slides.add_slide(BLANK)
top = header(s, 6, "Data sources and tools", "Everything we plan to use is free and open — no paid software or data")
data = [
    ("Component", "Tools / Data", "Remarks"),
    ("Satellite data", "Landsat 8/9 (thermal), Sentinel-2, MODIS — via Google Earth Engine", "Free; no heavy downloads needed"),
    ("Processing + ML", "Python, rasterio/GDAL, scikit-learn, XGBoost", "Standard open-source libraries"),
    ("Outputs", "Hotspot maps, ward tables, trend charts (PDF + web page)", "Simple formats planners already use"),
    ("Cross-checking", "IMD station data, Bhuvan maps, ground readings if available", "To verify our predictions"),
]
simple_table(s, 0.8, top + 0.1, 11.73, data, [2.3, 5.6, 3.83], size=12, row_h=0.68)
note_box(s, 0.8, 5.7, 11.73, 0.85, "Estimated cost of data and software: zero. The same method can be repeated for any Indian city.")
notes(s, "Main message: free stack, repeatable for any city. This handles objections about cost and feasibility.")

# ============ 7 · EXPECTED OUTCOMES (no dashboard!) ============
s = prs.slides.add_slide(BLANK)
top = header(s, 7, "What we expect to deliver", "Our planned outputs — with a sample trend to show the idea (illustrative values)")
# left: plain simple chart
cd = ChartData()
cd.categories = ["2015", "2017", "2019", "2021", "2023", "2025", "2030*"]
cd.add_series("Mean summer LST (°C)", (38.2, 39.1, 40.0, 40.8, 42.1, 43.0, 44.6))
cf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.8), Inches(top + 0.15), Inches(5.4), Inches(3.3), cd)
ch = cf.chart
ch.has_title = True
ch.chart_title.text_frame.text = "Sample LST trend (expected output format)"
for p in ch.chart_title.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(11); r.font.bold = True; r.font.name = FONT
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ser = ch.series[0]
ser.format.line.color.rgb = RED
ser.format.line.width = Pt(2.5)
try:
    ch.value_axis.minimum_scale = 36; ch.value_axis.maximum_scale = 46
except Exception:
    pass
txt(s, 0.8, top + 3.55, 5.4, 0.5, [[("*2030 is the model forecast. Values are samples to show the output format.", 9.5, LGRAY, False, None)]])
# right: deliverables
txt(s, 6.7, top + 0.1, 5.8, 0.4, [[("Planned deliverables:", 13, RED, True, None)]])
bullets(s, 6.7, top + 0.55, 5.8, [
    "Ward-wise 2030 heat-risk table for the pilot city.",
    "Hotspot map layers a planner can open in any GIS tool.",
    "Short action note per high-risk ward (plantation, cool roofs, etc.).",
    "Validation report comparing predictions with IMD records.",
], size=12.5, gap=8)
note_box(s, 6.7, 5.35, 5.83, 1.15, "Honest note: these are expected outputs. Real numbers will come only after we build the data pipeline.")
notes(s, "Be honest: this chart shows the OUTPUT FORMAT with sample numbers, not real results. Then read the four deliverables.")

# ============ 8 · SYSTEM ARCHITECTURE ============
s = prs.slides.add_slide(BLANK)
top = header(s, 8, "System architecture", "How the parts connect — satellite data in, planning decisions out")

def arch_band(y, label, sub, boxes, fillc):
    shape_rect(s, 0.8, y, 11.73, 1.28, fill=fillc, line=LINE, lw=0.75)
    shape_rect(s, 0.8, y, 0.07, 1.28, fill=RED)
    txt(s, 1.05, y + 0.24, 1.8, 0.9, [
        [(label, 11.5, RED, True, None)],
        [(sub, 9.5, GRAY, False, None)],
    ])
    bw = 2.18
    x = 3.05
    for b in boxes:
        shape_rect(s, x, y + 0.22, bw, 0.84, fill=WHITE, line=LGRAY, lw=0.9, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tx2 = s.shapes.add_textbox(Inches(x + 0.08), Inches(y + 0.24), Inches(bw - 0.16), Inches(0.8))
        tf2 = tx2.text_frame; tf2.word_wrap = True
        tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
        p0 = tf2.paragraphs[0]; p0.alignment = PP_ALIGN.CENTER
        r0 = p0.add_run(); r0.text = b
        r0.font.size = Pt(10.5); r0.font.bold = True; r0.font.color.rgb = INK; r0.font.name = FONT
        x += bw + 0.18

arch_band(top + 0.05, "LAYER 1", "Data sources",
          ["Landsat 8/9\nthermal (LST)", "Sentinel-2\nland-use images", "MODIS LST\n+ IMD records", "Ward / city\nboundary maps"], BOX)
txt(s, 6.3, top + 1.36, 0.7, 0.32, [[("▼", 13, LGRAY, True, None)]], align=PP_ALIGN.CENTER)
arch_band(top + 1.68, "LAYER 2", "Processing + model",
          ["Clean + align\ndata (GEE)", "Build features:\nNDVI, NDBI…", "Train ML model\n(RF / XGBoost)", "Forecast\n2030 hotspots"], WHITE)
txt(s, 6.3, top + 2.99, 0.7, 0.32, [[("▼", 13, LGRAY, True, None)]], align=PP_ALIGN.CENTER)
arch_band(top + 3.31, "LAYER 3", "Outputs + users",
          ["Hotspot maps\n(GIS layers)", "Ward heat-risk\ntables", "Ward action\nnotes", "City heat-plan\nreport"], BOX)
txt(s, 0.8, top + 4.72, 11.73, 0.35, [[("Build order:  data pipeline  →  model  →  outputs  →  field testing    ·    Pilot on one city first (about 7 weeks)", 10.5, GRAY, False, None)]])
notes(s, "Explain the three layers: what data comes in, what the model does, what the planner gets. Mention the build-order line at the bottom.")

# ============ 9 · IMPACT AND USERS ============
s = prs.slides.add_slide(BLANK)
top = header(s, 9, "Who will use this and why it matters", "Connecting our output to real decisions in Indian cities")
bullets(s, 0.8, top, 11.7, [
    ("Municipal corporations: ", "decide where to plant trees, protect lakes, and apply cool-roof rules — ward by ward."),
    ("Heat Action Plans (NDMA / state bodies): ", "get ward-level evidence to prepare city heat plans like Ahmedabad's."),
    ("Town planners and builders: ", "check the likely heat impact of new layouts before giving approvals."),
    ("Citizens: ", "cooler neighbourhoods, lower electricity bills, safer summers for workers and the elderly."),
    ("Scale: ", "50+ Indian cities face serious heat-island stress — the same method can serve all of them."),
], size=13.5, gap=9)
notes(s, "Name real users. End with scale: 50+ cities, same method, zero data cost. This is the impact slide — speak with energy.")

# ============ 10 · TEAM + THANK YOU ============
s = prs.slides.add_slide(BLANK)
top = header(s, 10, "Thank you!", "Our team — questions are welcome")
data = [
    ("Team member", "Role in this project"),
    ("[ Member 1 — Name ]", "Team lead — ML model and overall pipeline"),
    ("[ Member 2 — Name ]", "Satellite data and LST processing (GEE)"),
    ("[ Member 3 — Name ]", "Backend work — data handling and scripts"),
    ("[ Member 4 — Name ]", "Frontend/outputs — maps, charts and report"),
    ("[ Member 5 — Name ]", "Research — planning use-cases and validation"),
    ("[ Member 6 — Name ]", "Documentation, testing and presentation"),
]
simple_table(s, 0.8, top + 0.1, 11.73, data, [4.2, 7.53], size=12, row_h=0.5, first_col_bold=False)
txt(s, 0.8, 6.2, 11.73, 0.4, [[("Contact: [ email ]   |   GitHub: [ link ]   |   Mentor: [ name, if any ]", 12, GRAY, False, None)]])
notes(s, "Thank the judges, read the roles quickly, and invite questions. Keep the appendix ready.")

# ============ 11 · APPENDIX A ============
s = prs.slides.add_slide(BLANK)
top = header(s, 11, "Appendix A — LST calculation steps", "For reference, if asked how surface temperature is calculated")
steps_txt = [
    "Step 1 — Convert satellite DN values to radiance:   Lλ = ML × DN + AL",
    "Step 2 — Convert radiance to brightness temperature:   BT = K2 / ln(K1/Lλ + 1) − 273.15",
    "Step 3 — Estimate vegetation fraction from NDVI:   Pv = [(NDVI − NDVImin)/(NDVImax − NDVImin)]²",
    "Step 4 — Estimate surface emissivity:   ε = 0.004 × Pv + 0.986",
    "Step 5 — Final LST:   LST = BT / [ 1 + (λ·BT/ρ) · ln(ε) ]",
    "",
    "Then:  UHI intensity = average LST(urban wards) − average LST(rural surroundings)",
    "Trend per ward = yearly slope of summer LST from 2015 to 2025",
    "Forecast model:  LST(2030) = function of (past LST, NDVI, NDBI, built-up %) — checked using RMSE and R²",
]
tx = s.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.7), Inches(4.2))
tf = tx.text_frame; tf.word_wrap = True
first = True
for line in steps_txt:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(7)
    r = p.add_run(); r.text = line if line else " "
    r.font.size = Pt(12.5); r.font.color.rgb = INK; r.font.name = MONO
note_box(s, 0.8, 5.9, 11.73, 0.65, "Reference: Landsat Collection 2 handbook (USGS); Sobrino et al. NDVI-threshold emissivity method.")
notes(s, "Only show if asked. Walk through the five steps briefly, then the UHI and model lines.")

from pathlib import Path as _P
_OUT = _P(__file__).resolve().parent / "HeatSight_UHI_PPT_Human.pptx"
prs.save(str(_OUT))
print(f"Saved {_OUT}")
